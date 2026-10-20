#!/usr/bin/env python3
"""Convert all archive/canvas-submissions/*.docx to docs/<slug>.md."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parents[1]
ARCHIVE = ROOT / "archive" / "canvas-submissions"
DOCS = ROOT / "docs"


def slugify(stem: str) -> str:
    s = stem.lower()
    s = re.sub(r"[^a-z0-9]+", "-", s).strip("-")
    return s[:72] or "document"


def table_to_md(table) -> list[str]:
    rows = [[c.text.strip().replace("\n", " ") for c in row.cells] for row in table.rows]
    if not rows:
        return []
    lines = [
        "| " + " | ".join(rows[0]) + " |",
        "| " + " | ".join("---" for _ in rows[0]) + " |",
    ]
    for row in rows[1:]:
        lines.append("| " + " | ".join(row) + " |")
    return lines


def convert(src: Path) -> Path:
    doc = Document(src)
    rel = src.relative_to(ROOT).as_posix().replace(" ", "%20")
    title = src.stem
    lines: list[str] = [
        f"# {title}",
        "",
        f"**Source:** [`{src.name}`](../{rel})  ",
        "**Course:** IA 445 - Ethical Hacking and Offensive Security (EMU)  ",
        "**Author:** Julio Barros",
        "",
    ]
    for para in doc.paragraphs:
        t = para.text.strip()
        if not t:
            lines.append("")
            continue
        style = (para.style.name or "").lower()
        if "heading 1" in style:
            lines.extend(["", f"## {t}", ""])
        elif "heading 2" in style:
            lines.extend(["", f"### {t}", ""])
        elif "list" in style:
            lines.append(f"- {t}")
        else:
            lines.append(t)
            lines.append("")
    for table in doc.tables:
        lines.extend(["", *table_to_md(table), ""])
    body = re.sub(r"\n{3,}", "\n\n", "\n".join(lines)).strip() + "\n"
    dest = DOCS / f"{slugify(src.stem)}.md"
    dest.write_text(body, encoding="utf-8")
    return dest


def main() -> None:
    DOCS.mkdir(parents=True, exist_ok=True)
    for src in sorted(ARCHIVE.glob("*.docx")):
        if src.name.startswith("~$"):
            continue
        dest = convert(src)
        print(f"wrote {dest.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
